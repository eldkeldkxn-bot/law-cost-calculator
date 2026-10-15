"""간단한 소송비용계산서
설치: python -m pip install streamlit python-docx
실행: python -m streamlit run app.py
"""
from decimal import Decimal, ROUND_FLOOR
from io import BytesIO
import html
import re
import streamlit as st
from fractions import Fraction

# 2020.12.28. 개정 변호사보수 규칙 별표
BRACKETS = [
    (3_000_000, 0, 300_000, "0"),
    (20_000_000, 3_000_000, 300_000, ".10"),
    (50_000_000, 20_000_000, 2_000_000, ".08"),
    (100_000_000, 50_000_000, 4_400_000, ".06"),
    (150_000_000, 100_000_000, 7_400_000, ".04"),
    (200_000_000, 150_000_000, 9_400_000, ".02"),
    (500_000_000, 200_000_000, 10_400_000, ".01"),
    (None, 500_000_000, 13_400_000, ".005"),
]

STAGES = {
    # 민사
    "가합": 1, "가단": 1, "가소": 1, "나": 2, "다": 3,
    # 행정
    "구합": 1, "구단": 1, "누": 2, "두": 3,
    # 가사
    "드합": 1, "드단": 1, "르": 2, "므": 3
}
DELIVERY_UNIT = 5_640
LAW_FIRM = '법무법인(유한)바른'
DELIVERY_ROUNDS = {
    '가합': 15, '가단': 15, '가소': 10, '나': 12, '다': 8,
    '구합': 15, '구단': 15, '누': 12, '두': 8,
    '드합': 15, '드단': 15, '르': 12, '므': 8
}

SPAN_ITEMS = {'소계', '신청비용소계', '본안비용소계', '합계'}


def won(value):
    return f"{value:,}원"


def num_str(value):
    return f"{value:,}"


def trunc_10(value):
    """10원 미만 절사 (국고금 관리법 준용)"""
    return (int(value) // 10) * 10


def parse_fraction(text):
    text = text.strip()
    if not text:
        return Fraction(1, 1)
    try:
        if '/' in text:
            num, den = map(int, text.split('/'))
            return Fraction(num, den)
        else:
            return Fraction(int(text), 1)
    except ValueError:
        raise ValueError("부담비율은 '1/3', '1/1'과 같은 분수 형식으로 입력해주세요.")


def format_frac(frac):
    """분수를 법원 서식 비율(백분율 또는 기약분수)로 변환"""
    if frac == Fraction(1, 1):
        return "100%"
    if frac == Fraction(0, 1):
        return "0%"
    if 100 % frac.denominator == 0:
        return f"{int((frac.numerator / frac.denominator) * 100)}%"
    return f"{frac.numerator}/{frac.denominator}"


def fee_limit(soga):
    if soga <= 0:
        return 0
    for upper, lower, base, rate in BRACKETS:
        if upper is None or soga <= upper:
            value = Decimal(base) + Decimal(soga - lower) * Decimal(rate)
            return int(value.to_integral_value(rounding=ROUND_FLOOR))
    return 0


def fee_formula(soga, is_reduced=False):
    if soga <= 0:
        return "0원"
    for upper, lower, base, rate in BRACKETS:
        if upper is None or soga <= upper:
            if rate == '0':
                formula = won(base)
            else:
                percent = format(Decimal(rate) * 100, 'f').rstrip('0').rstrip('.')
                formula = f"{base:,}원 + ({soga:,}원 − {lower:,}원) × {percent}%"
            if is_reduced:
                formula = f"({formula}) × 1/2 (제5조 보수 감액)"
            return formula


def parse_cases(text):
    if not text.strip():
        return [{"stage": 1, "court": "", "number": "", "code": "가단"}]

    cases = []
    used = set()
    code_pattern = "|".join(STAGES.keys())
    regex = rf"(.*?)\s*(\d{{4}})\s*({code_pattern})\s*(\d+)"
    for line in re.split(r"[\n,;]+", text):
        if not line.strip():
            continue
        match = re.fullmatch(regex, line.strip())
        if not match:
            raise ValueError("사건번호 형식을 확인해주세요. (예: 서울중앙지방법원 2023가합12345)")
        court, year, code, number = match.groups()
        stage = STAGES[code]
        if stage in used:
            raise ValueError("같은 심급은 한 건씩 입력해주세요.")
        used.add(stage)
        cases.append({"stage": stage, "court": court.strip(), "number": year + code + number, "code": code})
    if not cases:
        return [{"stage": 1, "court": "", "number": "", "code": "가단"}]
    return sorted(cases, key=lambda case: case["stage"])


def optional_money(text):
    text = text.strip().replace(",", "").removesuffix("원").strip()
    if not text:
        return None
    if not re.fullmatch(r"[0-9]+", text):
        raise ValueError("금액은 0 이상의 숫자로 입력해주세요.")
    return int(text)


def filing_estimate(soga, case, electronic=True, party_count=1):
    if soga <= 0 or party_count < 1:
        return 0, 0
    if soga < 10_000_000:
        numerator = soga * 50
    elif soga < 100_000_000:
        numerator = soga * 45 + 50_000_000
    elif soga < 1_000_000_000:
        numerator = soga * 40 + 550_000_000
    else:
        numerator = soga * 35 + 5_550_000_000

    base = max(1_000, (numerator // 1_000_000) * 100)
    multiplier = {1: Decimal("1.0"), 2: Decimal("1.5"), 3: Decimal("2.0")}[case['stage']]
    stamp = int(Decimal(base) * multiplier)
    if electronic:
        stamp = stamp * 9 // 10
    stamp = max(1_000 if not electronic else 900, (stamp // 100) * 100)

    rounds = DELIVERY_ROUNDS.get(case['code'], 15)
    delivery = party_count * rounds * DELIVERY_UNIT
    return stamp, delivery


def make_model(cases, soga_dict, actual_fees, filing_costs, options):
    is_reduced = options.get('is_reduced', False)
    stage_settings = options.get('stage_settings', {})

    rows = []
    filing_rows = []
    main_total = 0
    borne_main_total = 0

    for case in cases:
        c_num = case["number"]
        c_key = c_num if c_num else f"stage_{case['stage']}"
        stage = case["stage"]
        soga = soga_dict.get(c_key, 0)
        if soga <= 0:
            raise ValueError(f"{stage}심 소가를 올바르게 입력해주세요.")

        limit = fee_limit(soga)
        if is_reduced:
            limit = limit // 2

        actual = actual_fees.get(c_key)
        if actual is not None and actual < 0:
            raise ValueError(f"{stage}심 실제 보수는 음수일 수 없습니다.")

        calc_fee = limit if actual is None else min(limit, actual)

        f_entry = filing_costs.get(c_key, {})
        include = f_entry.get('include', False)
        stamp = f_entry.get('stamp', 0) if include else 0
        delivery = f_entry.get('delivery', 0) if include else 0
        filing_rows.append(dict(stage=stage, number=c_num, stamp=stamp, delivery=delivery, include=include))

        settings = stage_settings.get(c_key, {'fraction': Fraction(1, 1), 'manual_name': '기타수기비용', 'manual_add': 0})
        fraction = settings['fraction']
        manual_name = settings['manual_name']
        manual_add = settings['manual_add']

        stage_sum = calc_fee + stamp + delivery + manual_add
        stage_borne = trunc_10(Decimal(stage_sum) * Decimal(fraction.numerator) / Decimal(fraction.denominator))

        rows.append({
            "stage": stage,
            "number": c_num,
            "soga": soga,
            "limit": limit,
            "actual": actual,
            "calc_fee": calc_fee,
            "stamp": stamp,
            "delivery": delivery,
            "manual_name": manual_name,
            "manual_add": manual_add,
            "fraction": fraction,
            "stage_sum": stage_sum,
            "stage_borne": stage_borne
        })

        main_total += stage_sum
        borne_main_total += stage_borne

    resp_count = options.get('respondent_count', 1)
    app_stamp = 900 if options.get('electronic', True) else 1_000
    app_delivery = (resp_count + 1) * 3 * DELIVERY_UNIT
    application_total = app_stamp + app_delivery

    final_total = borne_main_total + application_total
    provisional = any(r["actual"] is None for r in rows)
    is_all_won = all(r['fraction'] == Fraction(1, 1) for r in rows)

    # 상대방 소송비용 데이터 산출
    use_offset = options.get('use_offset', False)
    hide_opp_in_doc = options.get('hide_opp_in_doc', False)
    # 실제 문서(표/줄글/신청취지)에 피신청인 상계 내역을 반영할지 여부
    apply_offset_in_doc = use_offset and not hide_opp_in_doc

    opp_rows = []
    opp_total = 0
    opp_settings = options.get('opp_settings', {})

    for case in cases:
        c_num = case["number"]
        c_key = c_num if c_num else f"stage_{case['stage']}"
        stage = case["stage"]
        soga = soga_dict.get(c_key, 0)
        settings = stage_settings.get(c_key, {'fraction': Fraction(1, 1)})
        app_burden_fraction = max(Fraction(0, 1), Fraction(1, 1) - settings['fraction'])

        o_entry = opp_settings.get(c_key, {})
        o_limit = fee_limit(soga)
        if is_reduced:
            o_limit = o_limit // 2
        o_actual = o_entry.get('fee')
        o_calc_fee = o_limit if o_actual is None else min(o_limit, o_actual)

        o_stamp = o_entry.get('stamp', 0) if o_entry.get('include', False) else 0
        o_delivery = o_entry.get('delivery', 0) if o_entry.get('include', False) else 0
        o_manual = o_entry.get('manual', 0)

        o_sum = o_calc_fee + o_stamp + o_delivery + o_manual
        o_borne = trunc_10(Decimal(o_sum) * Decimal(app_burden_fraction.numerator) / Decimal(app_burden_fraction.denominator))

        opp_rows.append({
            "stage": stage,
            "number": c_num,
            "calc_fee": o_calc_fee,
            "actual": o_actual,
            "limit": o_limit,
            "stamp": o_stamp,
            "delivery": o_delivery,
            "manual": o_manual,
            "opp_sum": o_sum,
            "app_fraction": app_burden_fraction,
            "opp_borne": o_borne
        })
        if use_offset:
            opp_total += o_borne

    offset_final = final_total - opp_total if use_offset else final_total

    return dict(
        rows=rows, filing_rows=filing_rows, main_total=main_total,
        borne_main_total=borne_main_total, application_total=application_total,
        app_stamp=app_stamp, app_delivery=app_delivery, total=final_total,
        provisional=provisional, respondent_count=resp_count, is_all_won=is_all_won,
        use_offset=use_offset, hide_opp_in_doc=hide_opp_in_doc,
        apply_offset_in_doc=apply_offset_in_doc,
        opp_rows=opp_rows, opp_total=opp_total, offset_final=offset_final
    )


def application_sections(cases, model, options):
    refs = [f"{c['court']} {c['number']}".strip() for c in cases]
    refs = [r for r in refs if r]
    references = ', '.join(refs) if refs else '[법원 및 사건번호]'
    name = options['case_name'].strip() or '[사건명]'
    lawyer = LAW_FIRM
    date_text = (f"이 판결은 {options['final_date'].strip()} 확정되었습니다."
                 if options['final_date'].strip() else '')
    is_all_won = model['is_all_won']
    apply_offset = model.get('apply_offset_in_doc', False)
    claim_amount = model['offset_final'] if apply_offset and model['offset_final'] > 0 else model['total']

    prayer = (f'위 당사자 사이의 {references} {name} 사건의 판결에 의하여 피신청인이 신청인에게 '
              f'상환해야 할 소송비용액은 금 {won(claim_amount)}임을 확정한다.\n라는 결정을 구합니다.')

    if is_all_won:
        opening = f'신청인({options["role"]})을 상대로 피신청인이 제기한 {references} {name} 사건에서 최종적으로 전부승소하였습니다. ' if options['role'] == '피고' else f'신청인({options["role"]})이 피신청인을 상대로 제기한 {references} {name} 사건에서 최종적으로 전부승소하였습니다. '
        action = '응소' if options['role'] == '피고' else '제기·수행'
        burden_text = '위 사건의 소송비용은 피신청인이 전액 부담하도록 정하여졌으므로'
    else:
        opening = f'신청인({options["role"]})을 상대로 피신청인이 제기한 {references} {name} 사건에서 최종적으로 일부승소 판결이 선고되었습니다. ' if options['role'] == '피고' else f'신청인({options["role"]})이 피신청인을 상대로 제기한 {references} {name} 사건에서 최종적으로 일부승소 판결이 선고되었습니다. '
        action = '응소' if options['role'] == '피고' else '제기·수행'
        burden_text = '위 사건들의 각 판결에서 정해진 피신청인의 소송비용 부담비율을 적용하여 산정하였으므로'

    offset_mention = " 또한, 당사자 쌍방의 부담액에 관하여 민사소송법 제114조에 따라 대등액에서 상계하고 남은 잔액을 청구합니다." if apply_offset else ""

    cause = (
        f'1. {opening}{date_text}\n\n'
        f'2. 신청인은 위 사건을 {action}하기 위하여 {lawyer}을 소송대리인으로 선임하였으며, 상환을 구하는 비용은 별지 소송비용액계산서와 같습니다.\n\n'
        f'3. {burden_text},{offset_mention} 신청인은 별지와 같은 금원의 상환을 구하기 위하여 이 사건 신청에 이르렀습니다.'
    )
    return prayer, cause


def generate_table_data(cases, model, options):
    """표 데이터 생성 (피신청인 표시 여부에 따라 4열 또는 5열 호환)"""
    table_rows = []
    show_opp = model.get('apply_offset_in_doc', False)
    is_all_won = model.get('is_all_won', True)

    for idx, row in enumerate(model['rows']):
        stage_str = f"{row['stage']}심"
        stage_rows = []
        opp_row = model['opp_rows'][idx] if idx < len(model.get('opp_rows', [])) else {}
        formula_str = fee_formula(row['soga'], options.get('is_reduced', False))

        # 1. 변호사보수
        if row['actual'] is not None and row['actual'] < row['limit']:
            app_note = f"규칙 제3조 산정: {formula_str} = {won(row['limit'])}이나, 실제 지급한 변호사보수를 산입함"
        elif row['actual'] is not None:
            app_note = f"변호사보수의 소송비용산입에 관한 규칙 제3조에 의거 산정하면\n{formula_str} (실제 지급액: {won(row['actual'])})"
        else:
            app_note = f"변호사보수의 소송비용산입에 관한 규칙 제3조에 의거 산정하면\n{formula_str}"

        opp_val_str = '-'
        opp_note = ''
        if show_opp and opp_row:
            opp_val_str = won(opp_row['calc_fee'])
            if opp_row['actual'] is not None and opp_row['actual'] < opp_row['limit']:
                opp_note = f"규칙 제3조 산정: {formula_str} = {won(opp_row['limit'])}이나, 실제 지급한 변호사보수를 산입함"
            elif opp_row['actual'] is not None:
                opp_note = f"규칙 제3조 산정: {formula_str} = {won(opp_row['limit'])} (실제 지급액: {won(opp_row['actual'])})"
            else:
                opp_note = f"규칙 제3조 산정: {formula_str} = {won(opp_row['limit'])}"

        if show_opp:
            note_lines = [f"[신청인] {app_note}"]
            if opp_note:
                note_lines.append(f"[피신청인] {opp_note}")
        else:
            note_lines = [app_note]

        stage_rows.append({
            '비목': '변호사보수',
            'app_val': won(row['calc_fee']),
            'opp_val': opp_val_str,
            '비고': "\n".join(note_lines),
            'row_type': 'item'
        })

        # 2. 인지대
        has_app_stamp = row['stamp'] > 0
        has_opp_stamp = show_opp and opp_row.get('stamp', 0) > 0
        if has_app_stamp or has_opp_stamp:
            stage_rows.append({
                '비목': '인지대',
                'app_val': won(row['stamp']) if has_app_stamp else '-',
                'opp_val': won(opp_row['stamp']) if has_opp_stamp else '-',
                '비고': '',
                'row_type': 'item'
            })

        # 3. 송달료
        has_app_del = row['delivery'] > 0
        has_opp_del = show_opp and opp_row.get('delivery', 0) > 0
        if has_app_del or has_opp_del:
            stage_rows.append({
                '비목': '송달료',
                'app_val': won(row['delivery']) if has_app_del else '-',
                'opp_val': won(opp_row['delivery']) if has_opp_del else '-',
                '비고': '실제 지출비용 산입',
                'row_type': 'item'
            })

        # 4. 기타수기비용
        has_app_man = row['manual_add'] > 0
        has_opp_man = show_opp and opp_row.get('manual', 0) > 0
        if has_app_man or has_opp_man:
            m_name = row['manual_name'] if has_app_man else '기타 비용'
            stage_rows.append({
                '비목': m_name,
                'app_val': won(row['manual_add']) if has_app_man else '-',
                'opp_val': won(opp_row['manual']) if has_opp_man else '-',
                '비고': '',
                'row_type': 'item'
            })

        # 비목 행 개수 (법원 결정문처럼 '1심' 세로 병합은 비목 행들만 묶고 '소계'는 가로 2칸 병합)
        bimok_count = len(stage_rows)

        # 5. 소계
        stage_rows.append({
            '비목': '소계',
            'app_val': won(row['stage_sum']),
            'opp_val': won(opp_row['opp_sum']) if (show_opp and opp_row) else '-',
            '비고': '',
            'row_type': 'subtotal',
            'is_span_row': True
        })

        # 일부승소이거나 상계 표시일 때만 부담비율/부담비용 행 추가 (전부승소 단독일 때는 법원 결정문처럼 깔끔하게 생략 가능하되 명확성을 위해 비율이 100%가 아닐 때나 상계 시 표시)
        if not is_all_won or show_opp:
            app_ratio_str = format_frac(row['fraction'])
            opp_ratio_str = format_frac(opp_row.get('app_fraction', Fraction(0, 1))) if opp_row else '-'
            stage_rows.append({
                '비목': '상대방부담비율',
                'app_val': app_ratio_str,
                'opp_val': opp_ratio_str,
                '비고': '',
                'row_type': 'subtotal',
                'is_span_row': True
            })
            stage_rows.append({
                '비목': '상대방부담비용',
                'app_val': won(row['stage_borne']),
                'opp_val': won(opp_row['opp_borne']) if (show_opp and opp_row) else '-',
                '비고': '',
                'row_type': 'subtotal',
                'is_span_row': True
            })

        for i, sr in enumerate(stage_rows):
            is_span = sr.get('is_span_row', False)
            table_rows.append({
                '심급': stage_str if i == 0 else (sr['비목'] if is_span else ''),
                '비목': '' if is_span else sr['비목'],
                'app_val': sr['app_val'],
                'opp_val': sr['opp_val'],
                '비고': sr['비고'],
                'merge_len': bimok_count if i == 0 else 0,
                'is_span_row': is_span,
                'row_type': sr.get('row_type', 'item')
            })

    # 본안비용소계 행 (법원 결정문 양식에 맞춰 항상 표시)
    prefix_app = "[신청인] " if show_opp else ""
    note_sub = prefix_app + " + ".join([f"{r['stage']}심 {won(r['stage_borne'])}" for r in model['rows']])
    if show_opp:
        note_sub += "\n[피신청인] " + " + ".join([f"{r['stage']}심 {won(r['opp_borne'])}" for r in model['opp_rows']])
    table_rows.append({
        '심급': '본안비용소계',
        '비목': '',
        'app_val': won(model['borne_main_total']),
        'opp_val': won(model['opp_total']) if show_opp else '-',
        '비고': note_sub,
        'merge_len': 0,
        'is_span_row': True,
        'row_type': 'main_subtotal'
    })

    # 소송비용확정신청 행 (인지대, 송달료 2행만 세로 병합하고 신청비용소계는 가로 2칸 병합)
    app_bimok_rows = [
        {'비목': '인지대', 'app_val': won(model['app_stamp']), 'opp_val': '-', '비고': ''},
        {'비목': '송달료', 'app_val': won(model['app_delivery']), 'opp_val': '-', '비고': f"피신청인 {model['respondent_count']}명 기준 (잔액은 확정 후 환급)"},
    ]
    for i, ar in enumerate(app_bimok_rows):
        table_rows.append({
            '심급': '소송비용\n확정신청' if i == 0 else '',
            '비목': ar['비목'],
            'app_val': ar['app_val'],
            'opp_val': ar['opp_val'],
            '비고': ar['비고'],
            'merge_len': len(app_bimok_rows) if i == 0 else 0,
            'is_span_row': False,
            'row_type': 'item'
        })

    # 신청비용소계
    table_rows.append({
        '심급': '신청비용소계',
        '비목': '',
        'app_val': won(model['application_total']),
        'opp_val': '-',
        '비고': '',
        'merge_len': 0,
        'is_span_row': True,
        'row_type': 'subtotal'
    })

    # 최종 합계 행
    if show_opp:
        total_note = f"[신청인] 본안비용 {won(model['borne_main_total'])} + 신청비용 {won(model['application_total'])}"
        total_note += f"\n[피신청인] 본안비용 {won(model['opp_total'])}"
        if model['offset_final'] >= 0:
            total_note += f"\n▶ 대등액 상계 후 최종 상환액: {won(model['offset_final'])}"
        else:
            total_note += f"\n▶ 대등액 상계 결과 신청인 초과부담: {won(abs(model['offset_final']))}"
    else:
        total_note = f"본안비용 {won(model['borne_main_total'])} + 신청비용 {won(model['application_total'])}"

    table_rows.append({
        '심급': '합계',
        '비목': '',
        'app_val': won(model['total']),
        'opp_val': won(model['opp_total']) if show_opp else '-',
        '비고': total_note,
        'merge_len': 0,
        'is_span_row': True,
        'row_type': 'grand_total'
    })

    return table_rows


def render_html_table(cases, model, options):
    """법원 결정문 스타일의 HTML 표 생성 (들여쓰기 코드블록 오작동 방지 적용)"""
    table_data = generate_table_data(cases, model, options)
    show_opp = model.get('apply_offset_in_doc', False)

    refs = [f"{c['court']} {c['number']}".strip() for c in cases if (c['court'] or c['number'])]
    ref_str = ", ".join(refs) if refs else "본안사건"
    soga_parts = [f"제{r['stage']}심 소가 : {won(r['soga'])}" for r in model['rows']]
    soga_line = " / ".join(soga_parts)

    burden_lines = []
    for r in model['rows']:
        stage_label = {1: "소송비용", 2: "항소비용", 3: "상고비용"}.get(r['stage'], f"{r['stage']}심 비용")
        if r['fraction'] == Fraction(1, 1):
            burden_lines.append(f"▶ {stage_label}은 피신청인({'원고' if options['role']=='피고' else '피고'})이 부담")
        else:
            burden_lines.append(f"▶ {stage_label} 중 {format_frac(r['fraction'])}은 피신청인이 부담")
    burden_html = "<br>".join(burden_lines)

    if show_opp:
        colgroup_html = (
            '<colgroup>'
            '<col style="width: 11%;">'
            '<col style="width: 14%;">'
            '<col style="width: 14%;">'
            '<col style="width: 14%;">'
            '<col style="width: 47%;">'
            '</colgroup>'
        )
        thead_html = (
            '<thead>'
            '<tr><th rowspan="2">심급</th><th rowspan="2">비목</th><th colspan="2">비용액</th><th rowspan="2">비고</th></tr>'
            '<tr><th>신청인</th><th>피신청인</th></tr>'
            '</thead>'
        )
    else:
        colgroup_html = (
            '<colgroup>'
            '<col style="width: 12%;">'
            '<col style="width: 16%;">'
            '<col style="width: 18%;">'
            '<col style="width: 54%;">'
            '</colgroup>'
        )
        thead_html = (
            '<thead>'
            '<tr><th>심급</th><th>비목</th><th>비용액</th><th>비고</th></tr>'
            '</thead>'
        )

    tbody_rows = []
    for row in table_data:
        r_type = row.get('row_type', 'item')
        tr_class = {
            'item': '',
            'subtotal': ' class="row-subtotal"',
            'main_subtotal': ' class="row-main-subtotal"',
            'grand_total': ' class="row-grand-total"'
        }.get(r_type, '')

        tds = []
        if row.get('is_span_row'):
            label_html = html.escape(row['심급']).replace('\n', '<br>')
            tds.append(f'<td colspan="2" class="cell-center">{label_html}</td>')
        else:
            if row.get('merge_len', 0) > 0:
                stage_html = html.escape(row['심급']).replace('\n', '<br>')
                rowspan_attr = f' rowspan="{row["merge_len"]}"' if row['merge_len'] > 1 else ''
                tds.append(f'<td{rowspan_attr} class="cell-stage">{stage_html}</td>')
            bimok_html = html.escape(row['비목']).replace('\n', '<br>')
            tds.append(f'<td class="cell-center">{bimok_html}</td>')

        app_v = str(row['app_val'])
        app_cls = "cell-right" if app_v.endswith("원") else "cell-center"
        tds.append(f'<td class="{app_cls}">{html.escape(app_v)}</td>')

        if show_opp:
            opp_v = str(row['opp_val'])
            opp_cls = "cell-right" if opp_v.endswith("원") else "cell-center"
            tds.append(f'<td class="{opp_cls}">{html.escape(opp_v)}</td>')

        note_html = html.escape(row['비고']).replace('\n', '<br>')
        tds.append(f'<td class="cell-left">{note_html}</td>')

        tbody_rows.append(f'<tr{tr_class}>{"".join(tds)}</tr>')

    claim_amount = model['offset_final'] if show_opp and model['offset_final'] > 0 else model['total']

    full_html = (
        '<div class="court-doc-card">'
        '<div class="court-doc-title">소 송 비 용 계 산 서</div>'
        '<div class="court-meta-box">'
        f'<div><strong>신청인 :</strong> {html.escape(options["role"])} ({html.escape(ref_str)})</div>'
        f'<div><strong>{html.escape(soga_line)}</strong></div>'
        f'<div class="court-burden-list">{burden_html}</div>'
        '</div>'
        '<div class="court-table-topbar">'
        '<span class="sec-label">1. 총 소송비용</span>'
        '<span class="unit-label">(단위 : 원, 10원 미만 버림)</span>'
        '</div>'
        f'<table class="court-cost-table">{colgroup_html}{thead_html}<tbody>{"".join(tbody_rows)}</tbody></table>'
        '<div class="court-reimburse-box">'
        f'2. 피신청인이 신청인에게 상환하여야 할 소송비용액 금 {won(claim_amount)}'
        '</div>'
        '</div>'
    )
    return full_html


def render_text_calculation(cases, model, options):
    show_opp = model.get('apply_offset_in_doc', False)
    before = ['1. 신청인의 지출비용 및 피신청인 부담액']
    for index, row in enumerate(model['rows']):
        prefix = '가나다라마바사'[index]
        num_str = f" ({row['number']})" if row['number'] else ""
        stage_title = f"{prefix}. {row['stage']}심{num_str}"
        before.append(f"{stage_title} 소가: 금 {won(row['soga'])}")

        actual_str = f" (실제 지급액: 금 {won(row['actual'])})" if row['actual'] is not None else ""
        before.append(f"   - 변호사보수: 금 {won(row['calc_fee'])}{actual_str}")
        before.append(f"   - 최대인정보수 산식: {fee_formula(row['soga'], options.get('is_reduced', False))}")

        if row['stamp']:
            before.append(f"   - 인지대: 금 {won(row['stamp'])}")
        if row['delivery']:
            before.append(f"   - 송달료: 금 {won(row['delivery'])}")
        if row['manual_add'] > 0:
            before.append(f"   - {row['manual_name']}: 금 {won(row['manual_add'])}")

        before.append(f"   => {row['stage']}심 지출 합계: 금 {won(row['stage_sum'])}")

        frac = row['fraction']
        if frac == Fraction(1, 1):
            before.append(f"   => 피신청인 부담비율(전부) 적용: 금 {won(row['stage_borne'])}")
        else:
            before.append(f"   => 피신청인 부담비율({frac.numerator}/{frac.denominator}) 적용: 금 {won(row['stage_borne'])}")
        before.append("")

    after = ['2. 소송비용액 산정']
    after.append(f"가. 본안 소송비용 중 피신청인 부담액 합계: 금 {won(model['borne_main_total'])}")
    after.append(f"나. 소송비용액확정신청 인지대 금 {won(model['app_stamp'])} 및 송달료 금 {won(model['app_delivery'])} (피신청인 {model['respondent_count']}명 기준)")
    after.append(f"다. 본안 및 신청비용 합계: {won(model['borne_main_total'])} + {won(model['application_total'])} = 금 {won(model['total'])}")

    if show_opp:
        after.append("")
        after.append("3. 상대방 소송비용과의 대등액 상계 (민사소송법 제114조)")
        after.append(f"가. 상대방 본안 소송비용 중 신청인 부담액 합계: 금 {won(model['opp_total'])}")
        if model['offset_final'] >= 0:
            after.append(f"나. 대등액 상계 후 피신청인이 상환해야 할 최종 잔액: 금 {won(model['total'])} − 금 {won(model['opp_total'])} = 금 {won(model['offset_final'])}   끝.")
        else:
            after.append(f"나. 대등액 상계 결과 (신청인의 초과 부담액 발생): 금 {won(abs(model['offset_final']))}   끝.")
    else:
        after.append("라. 피신청인이 신청인에게 상환해야 할 소송비용액은,")
        after.append(f"    금 {won(model['total'])}   끝.")

    return before, after


def make_docx(cases, model, options, prayer, cause):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_ALIGN_VERTICAL

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(3.5)

    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)

    for name, font, size in [('Normal', '나눔명조', 11.5), ('Title', '나눔고딕', 13)]:
        style = doc.styles[name]
        style.font.name = font
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.bold = True

    def heading(text):
        p = doc.add_paragraph(text, 'Title')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True

    def paragraphs(lines):
        for line in lines:
            if line.strip():
                p = doc.add_paragraph(line)
                for run in p.runs:
                    run.font.name = '나눔명조'
                    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '나눔명조')

    heading('신 청 취 지')
    paragraphs(prayer.split('\n'))
    doc.add_paragraph()
    heading('신 청 원 인')
    paragraphs(cause.split('\n'))
    doc.add_page_break()

    doc.add_paragraph('별지')
    heading('소 송 비 용 계 산 서')

    use_table_format = options.get('output_format') == '표 양식 (법원 서식)'
    show_opp = model.get('apply_offset_in_doc', False)

    if use_table_format:
        table_data = generate_table_data(cases, model, options)
        num_cols = 5 if show_opp else 4
        hdr_rows_count = 2 if show_opp else 1

        table = doc.add_table(rows=len(table_data) + hdr_rows_count, cols=num_cols)
        table.style = 'Table Grid'
        table.autofit = False

        tblPr = table._tbl.tblPr
        cell_mar = parse_xml(
            f'<w:tblCellMar {nsdecls("w")}>'
            f'<w:top w:w="110" w:type="dxa"/>'
            f'<w:bottom w:w="110" w:type="dxa"/>'
            f'<w:left w:w="140" w:type="dxa"/>'
            f'<w:right w:w="140" w:type="dxa"/>'
            f'</w:tblCellMar>'
        )
        tblPr.append(cell_mar)

        if show_opp:
            col_widths = [Cm(2.1), Cm(2.5), Cm(2.7), Cm(2.7), Cm(6.0)]
            # 2단 헤더 구성 (비용액 -> 신청인 | 피신청인)
            r0 = table.rows[0].cells
            r1 = table.rows[1].cells
            r0[0].text = '심급'
            r0[1].text = '비목'
            r0[2].text = '비용액'
            r0[4].text = '비고'
            r1[2].text = '신청인'
            r1[3].text = '피신청인'

            r0[0].merge(r1[0])
            r0[1].merge(r1[1])
            r0[2].merge(r0[3])
            r0[4].merge(r1[4])

            for h_idx in range(hdr_rows_count):
                for c in table.rows[h_idx].cells:
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
                    c._tc.get_or_add_tcPr().append(shd)
        else:
            col_widths = [Cm(2.2), Cm(2.8), Cm(3.0), Cm(8.0)]
            headers = ['심급', '비목', '비용액', '비고']
            for i, text in enumerate(headers):
                cell = table.rows[0].cells[i]
                cell.text = text
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
                cell._tc.get_or_add_tcPr().append(shd)

        # 본문 데이터 입력
        for r_idx, row in enumerate(table_data):
            t_row_idx = r_idx + hdr_rows_count
            cells = table.rows[t_row_idx].cells
            cells[0].text = row['심급']
            cells[1].text = row['비목']
            cells[2].text = str(row['app_val'])
            if show_opp:
                cells[3].text = str(row['opp_val'])
                cells[4].text = row['비고']
            else:
                cells[3].text = row['비고']

        # 셀 병합 및 음영 처리
        for r_idx, row in enumerate(table_data):
            t_row_idx = r_idx + hdr_rows_count
            m_len = row.get('merge_len', 0)
            if m_len > 1:
                start_cell = table.cell(t_row_idx, 0)
                end_cell = table.cell(t_row_idx + m_len - 1, 0)
                merged = start_cell.merge(end_cell)
                while len(merged.paragraphs) > 1:
                    p = merged.paragraphs[-1]._p
                    p.getparent().remove(p)

            if row.get('is_span_row'):
                start_cell = table.cell(t_row_idx, 0)
                end_cell = table.cell(t_row_idx, 1)
                merged = start_cell.merge(end_cell)
                while len(merged.paragraphs) > 1:
                    p = merged.paragraphs[-1]._p
                    p.getparent().remove(p)
                merged.paragraphs[0].text = row['심급']

                r_type = row.get('row_type', 'subtotal')
                fill_hex = "EFF6FF" if r_type == 'grand_total' else ("F1F5F9" if r_type == 'main_subtotal' else "F8FAFC")
                for c_i in range(num_cols):
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
                    table.cell(t_row_idx, c_i)._tc.get_or_add_tcPr().append(shd)

        # 폰트 및 정렬 일괄 적용
        for r_idx, row_obj in enumerate(table.rows):
            is_header = r_idx < hdr_rows_count
            data_row = table_data[r_idx - hdr_rows_count] if not is_header else None
            is_bold_row = is_header or (data_row and data_row.get('is_span_row'))

            for c_idx, cell in enumerate(row_obj.cells):
                cell.width = col_widths[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15

                    for run in p.runs:
                        run.font.name = '나눔명조'
                        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '나눔명조')
                        run.font.size = Pt(9.5)
                        if is_bold_row or c_idx == 0:
                            run.font.bold = True

                    if is_header:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        amt_cols = [2, 3] if show_opp else [2]
                        note_col = 4 if show_opp else 3
                        if c_idx in amt_cols:
                            val_text = p.text.strip()
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if val_text.endswith('원') else WD_ALIGN_PARAGRAPH.CENTER
                        elif c_idx == note_col:
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        else:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        before, after = render_text_calculation(cases, model, options)
        paragraphs(before)
        doc.add_paragraph()
        paragraphs(after)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def main():
    st.set_page_config(page_title='소송비용 신청서 만들기', page_icon='⚖️', layout='wide')

    st.markdown("""
    <style>
    .block-container {
        max-width: 980px !important;
        padding-top: 2rem !important;
    }
    [data-testid="stDownloadButton"] button {
        background-color: #1e3a8a;
        color: #ffffff;
        border: 1px solid #1e3a8a;
        font-weight: 700;
        padding: 0.65rem 1rem;
        border-radius: 6px;
    }
    [data-testid="stDownloadButton"] button:hover {
        background-color: #172554;
        color: #ffffff;
    }

    /* 법원 결정문 스타일 카드 & 테이블 */
    .court-doc-card {
        background: #ffffff;
        border: 1px solid #cbd5e1;
        border-radius: 8px;
        padding: 32px 36px;
        box-shadow: 0 4px 14px rgba(15, 23, 42, 0.05);
        margin-top: 12px;
        margin-bottom: 24px;
        color: #0f172a;
        font-family: 'Pretendard', 'Malgun Gothic', '맑은 고딕', serif;
    }
    .court-doc-title {
        text-align: center;
        font-size: 22px;
        font-weight: 800;
        letter-spacing: 0.25em;
        margin-bottom: 22px;
        color: #0f172a;
    }
    .court-meta-box {
        background: #f8fafc;
        border-left: 4px solid #334155;
        padding: 12px 16px;
        font-size: 14px;
        line-height: 1.6;
        margin-bottom: 16px;
        border-radius: 0 6px 6px 0;
    }
    .court-burden-list {
        margin-top: 6px;
        padding-top: 6px;
        border-top: 1px dashed #cbd5e1;
        color: #334155;
        font-size: 13.5px;
    }
    .court-table-topbar {
        display: flex;
        justify-content: space-between;
        align-items: flex-end;
        margin-bottom: 6px;
        font-size: 13.5px;
    }
    .court-table-topbar .sec-label {
        font-weight: 700;
        font-size: 14.5px;
    }
    .court-table-topbar .unit-label {
        color: #64748b;
        font-size: 12.5px;
    }
    .court-cost-table {
        width: 100% !important;
        border-collapse: collapse !important;
        border-top: 2px solid #0f172a !important;
        border-bottom: 2px solid #0f172a !important;
        font-size: 13.5px !important;
        table-layout: fixed !important;
        background: #ffffff !important;
    }
    .court-cost-table th,
    .court-cost-table td {
        border: 1px solid #cbd5e1 !important;
        padding: 9px 12px !important;
        vertical-align: middle !important;
        word-break: keep-all !important;
        line-height: 1.45 !important;
    }
    .court-cost-table thead th {
        background-color: #f1f5f9 !important;
        color: #0f172a !important;
        font-weight: 700 !important;
        text-align: center !important;
        border-bottom: 1.5px solid #64748b !important;
    }
    .cell-stage {
        background-color: #fcfcfd !important;
        font-weight: 700 !important;
        text-align: center !important;
        color: #0f172a !important;
    }
    .cell-center {
        text-align: center !important;
    }
    .cell-right {
        text-align: right !important;
        font-variant-numeric: tabular-nums !important;
        font-weight: 500 !important;
        white-space: nowrap !important;
    }
    .cell-left {
        text-align: left !important;
        color: #334155 !important;
        font-size: 13px !important;
    }
    .row-subtotal td {
        background-color: #f8fafc !important;
        font-weight: 600 !important;
    }
    .row-main-subtotal td {
        background-color: #f1f5f9 !important;
        font-weight: 700 !important;
        color: #0f172a !important;
        border-top: 1.5px solid #94a3b8 !important;
        border-bottom: 1.5px solid #94a3b8 !important;
    }
    .row-grand-total td {
        background-color: #eff6ff !important;
        font-weight: 700 !important;
        color: #1e3a8a !important;
        font-size: 14.5px !important;
        border-top: 2px solid #2563eb !important;
    }
    .court-reimburse-box {
        margin-top: 14px;
        padding: 12px 16px;
        background: #f8fafc;
        border: 1px solid #cbd5e1;
        border-radius: 6px;
        font-weight: 700;
        font-size: 14.5px;
        color: #0f172a;
    }
    </style>
    """, unsafe_allow_html=True)

    st.title('소송비용 신청서 만들기')
    st.caption('소송비용액확정 신청서 및 계산서 간편 생성기')

    case_column, name_column = st.columns([3, 1])
    with case_column:
        raw = st.text_area('1·2·3심 법원·사건번호',
                           placeholder='서울중앙지방법원 2023가합12345, 서울고등법원 2025나12345', height=100)
    with name_column:
        case_name = st.text_input('사건명', placeholder='약정금')
    st.caption('진행한 심급만 쉼표나 줄바꿈으로 구분하세요. 행정, 가사 사건도 입력 가능합니다.')

    error, cases = None, []
    try:
        cases = parse_cases(raw)
    except ValueError as exc:
        error = str(exc)
        cases = [{"stage": 1, "court": "", "number": "", "code": "가단"}]

    role = st.radio('신청인은 본안에서', ['피고', '원고'], horizontal=True)

    method_column, people_column = st.columns(2)
    with method_column:
        electronic = st.radio('제출 방식', ['전자소송', '종이소송'], horizontal=True) == '전자소송'
    with people_column:
        party_count = int(st.number_input('송달료 산정 기준 인원수', min_value=1,
                                          value=1 if electronic else 2, step=1))

    st.markdown("---")
    st.subheader("소송비용액확정 신청 비용 설정")
    resp_col1, resp_col2 = st.columns([1, 2])
    with resp_col1:
        respondent_count = int(st.number_input('상대방(피신청인) 수', min_value=1, value=1, step=1))
    with resp_col2:
        auto_app_stamp = 900 if electronic else 1_000
        total_parties = respondent_count + 1
        auto_rounds = total_parties * 3
        auto_app_delivery = auto_rounds * DELIVERY_UNIT
        st.info(f"📌 **신청 비용 자동 계산 내역**\n\n- 인지대: **{won(auto_app_stamp)}** (전자 기준)\n- 송달료: **{won(auto_app_delivery)}**")

    is_reduced = st.checkbox('무변론 판결, 자백간주, 이행권고결정에 따른 변호사보수 1/2 감액 일괄 적용 (규칙 제5조)')

    st.markdown("---")
    st.subheader("심급별 세부 비용 및 부담비율 설정")
    soga_dict = {}
    actual_fees = {}
    filing_costs = {}
    stage_settings = {}

    for case in cases:
        c_num = case['number']
        c_key = c_num if c_num else f"stage_{case['stage']}"
        stage = case['stage']
        court_info = f" ({case['court']} {c_num})".strip() if (case['court'] or c_num) else ""
        st.markdown(f"#### 🏛️ {stage}심{court_info}")

        soga_col, fee_col = st.columns(2)
        with soga_col:
            soga_val = int(st.number_input(f'{stage}심 소가(원)', min_value=0, value=50_000_000, step=1_000_000, key=f'soga_{c_key}'))
            soga_dict[c_key] = soga_val
        with fee_col:
            fee_val = st.text_input(f'{stage}심 실제 변호사보수(원)', key=f'fee_{c_key}',
                                    placeholder='생략 시 법정 한도액 적용')
            try:
                actual_fees[c_key] = optional_money(fee_val)
            except ValueError as exc:
                error = str(exc)

            col1, col2, col3 = st.columns([1.2, 1.5, 1.5])
            with col1:
                frac_str = st.text_input(f'{stage}심 피신청인 부담비율', value='1/1', key=f'frac_{c_key}', help="분수 입력")
                try:
                    stage_fraction = parse_fraction(frac_str)
                except ValueError as exc:
                    error = str(exc)
                    stage_fraction = Fraction(1, 1)
            with col2:
                manual_name = st.text_input(f'{stage}심 수기 비용 항목명', value='기타 수기 비용', key=f'manual_name_{c_key}')
            with col3:
                manual_add = int(st.number_input(f'{stage}심 기타 비용 금액(원)', min_value=0, value=0, step=10_000, key=f'manual_{c_key}', help="0원 입력시 계산서 적용 안됨"))

            stage_settings[c_key] = {'fraction': stage_fraction, 'manual_name': manual_name, 'manual_add': manual_add}

            include = st.checkbox(f'신청인이 납부한 {stage}심 인지대·송달료 포함',
                                  value=(role == '원고' and stage == 1),
                                  key=f'paid_{role}_{c_key}')
            stamp, delivery = (0, 0)
            if soga_val > 0:
                stamp, delivery = filing_estimate(soga_val, case, electronic, party_count)
            entry = dict(include=include, stamp=stamp, delivery=delivery)
            if include:
                st.caption(f"자동 계산: 인지대 {won(stamp)} / 송달료 {won(delivery)}")
                if st.checkbox(f'{stage}심 인지·송달료 실제 납부액으로 직접 수정', key=f'override_{c_key}'):
                    col_st, col_dl = st.columns(2)
                    with col_st:
                        entry['stamp'] = int(st.number_input(f'{stage}심 인지대 실부담액', min_value=0, value=stamp, key=f'stamp_{c_key}'))
                    with col_dl:
                        entry['delivery'] = int(st.number_input(f'{stage}심 송달료 실부담액', min_value=0, value=delivery, key=f'delivery_{c_key}'))
            filing_costs[c_key] = entry
            st.divider()

    st.caption('소송대리인: 법무법인(유한)바른')
    opp_settings = {}
    hide_opp_in_doc = False
    with st.expander('⚖️ 상대방 소송비용 상계(대등액 공제) 시뮬레이션 (선택)'):
        use_offset = st.checkbox('상대방 소송비용 상계 계산 활성화', value=False)
        if use_offset:
            hide_opp_in_doc = st.checkbox(
                '📄 신청서·계산서(표 및 줄글) 출력 시 피신청인(상대방) 내역 숨기기 (신청인 지출액만 단독 청구)',
                value=False,
                help="체크하면 상단에서 상계 시뮬레이션 금액은 확인하되, 실제 출력되는 신청서·표·줄글에서는 피신청인 칸과 상계 문구를 빼고 신청인 지출액 기준으로만 깔끔하게 생성합니다."
            )
            is_all_won = all(stage_settings.get(c['number'] or f"stage_{c['stage']}", {}).get('fraction', Fraction(1, 1)) == Fraction(1, 1) for c in cases)
            if is_all_won:
                st.warning("⚠️ **현재 피신청인 부담비율이 '1/1(전부승소)'입니다.**\n\n신청인이 물어줘야 할 비용이 0%이므로 상계액이 계산되지 않습니다. 부담비율을 수정해주세요.")

            for case in cases:
                c_num = case['number']
                c_key = c_num if c_num else f"stage_{case['stage']}"
                stage = case['stage']
                soga_val = soga_dict.get(c_key, 0)
                court_info = f" ({case['court']} {c_num})".strip() if (case['court'] or c_num) else ""
                st.markdown(f"**🏛️ {stage}심 상대방 지출비용{court_info}**")

                col_of1, col_of2 = st.columns(2)
                with col_of1:
                    opp_fee_text = st.text_input(f'{stage}심 상대방 실제 변호사보수(원)', key=f'opp_fee_{c_key}', placeholder='생략 시 법정 한도액 적용')
                    try:
                        opp_fee_val = optional_money(opp_fee_text)
                    except ValueError as exc:
                        error = str(exc)
                        opp_fee_val = None
                with col_of2:
                    opp_manual_val = int(st.number_input(f'{stage}심 상대방 기타 비용(원)', min_value=0, value=0, step=10_000, key=f'opp_man_{c_key}', help="0원 입력시 계산서 적용 안됨"))

                opp_include = st.checkbox(f'상대방이 납부한 {stage}심 인지대·송달료 포함',
                                          value=(role == '피고' and stage == 1),
                                          key=f'opp_paid_{c_key}')
                o_stamp, o_delivery = 0, 0
                if soga_val > 0:
                    o_stamp, o_delivery = filing_estimate(soga_val, case, electronic, party_count)
                if opp_include:
                    st.caption(f"상대방 인지·송달료 추정치: 인지대 {won(o_stamp)} / 송달료 {won(o_delivery)}")

                opp_settings[c_key] = {
                    'fee': opp_fee_val,
                    'include': opp_include,
                    'stamp': o_stamp,
                    'delivery': o_delivery,
                    'manual': opp_manual_val
                }
                st.markdown("---")

    with st.expander('확정일 입력 (선택)'):
        final_date = st.text_input('확정일', placeholder='예: 2026. 7. 21.')
        st.caption('입력하지 않으면 확정일 문장은 문서에 넣지 않습니다.')

    st.markdown("---")
    output_format = st.radio('📝 계산서 출력 양식', ['표 양식 (법원 서식)', '줄글 양식 (기본)'], horizontal=True)

    options = dict(electronic=electronic, party_count=party_count,
                   role=role, case_name=case_name, final_date=final_date,
                   is_reduced=is_reduced, respondent_count=respondent_count,
                   stage_settings=stage_settings, use_offset=use_offset,
                   hide_opp_in_doc=hide_opp_in_doc, opp_settings=opp_settings,
                   output_format=output_format)

    signature = repr((cases, raw, soga_dict, actual_fees, filing_costs, options))

    if st.button('신청서·계산서 만들기', type='primary', use_container_width=True):
        try:
            if error:
                raise ValueError(error)
            model = make_model(cases, soga_dict, actual_fees, filing_costs, options)
        except ValueError as exc:
            st.error(str(exc))
            st.session_state.pop('template_result', None)
        else:
            st.session_state['template_result'] = (signature, model)

    saved = st.session_state.get('template_result')
    if not saved or saved[0] != signature:
        return

    model = saved[1]
    st.success("✅ 계산서 작성이 완료되었습니다.")

    if model.get('use_offset'):
        col_m1, col_m2 = st.columns(2)
        with col_m1:
            st.metric('상계 전 신청인 단독 청구액', won(model['total']))
        with col_m2:
            delta_str = f"-{won(model['opp_total'])}" if model['opp_total'] > 0 else "0원"
            label_suffix = " (참고용·문서 미반영)" if model.get('hide_opp_in_doc') else ""
            st.metric(f'상계 후 예상 상환액{label_suffix}', won(model['offset_final']), delta=f"상대방 채권 공제: {delta_str}")
    else:
        st.metric('최종 상환액 합계', won(model['total']))

    prayer, cause = application_sections(cases, model, options)
    prayer = st.text_area('신청취지 편집', value=prayer, height=140)
    cause = st.text_area('신청원인 편집', value=cause, height=260)

    st.subheader('별지 소송비용액계산서')

    if output_format == '표 양식 (법원 서식)':
        html_table = render_html_table(cases, model, options)
        st.markdown(html_table, unsafe_allow_html=True)
    else:
        before, after = render_text_calculation(cases, model, options)
        for line in before:
            st.write(line)
        for line in after:
            st.write(line)

    args = (cases, model, options, prayer, cause)
    try:
        document = make_docx(*args)
    except ImportError:
        st.info('Word 저장: 터미널에서 python -m pip install python-docx를 실행하세요.')
    else:
        st.download_button('📥 신청서·계산서 초안 다운로드 (.docx)', document,
                           file_name='소송비용액확정_신청서_계산서.docx',
                           mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                           use_container_width=True)


if __name__ == '__main__':
    main()
